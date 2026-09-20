from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT = Path("docs/customer-validation-forms")
SITES = {
    "okapiccf.com": "Slow response snapshot; earlier security-header gaps.",
    "ritssbeauty.com": "No non-empty H1; HSTS/CSP gaps.",
    "hmcng.com": "No non-empty H1; HSTS/CSP/nosniff/frame-protection gaps.",
    "gabofarms.com": "Recent audit confirmed SEO content; security probes may be incomplete because of verification pages.",
    "gabogreenenergysolutions.com": "Sitemap availability and intermittent reachability require owner confirmation.",
    "mvpterminallingllc.com": "No meta description; HSTS/CSP/nosniff/frame-protection gaps.",
    "fareharbor.com": "Sitemap returned a redirect; CSP/frame-protection gaps; some probes incomplete.",
    "danwebdevelopment.com": "Missing meta description, H1, canonical URL, and sitemap; security-header gaps.",
    "nenafrika.com": "Slow response observed in latest audit; security-header gaps.",
    "africancouncilofoptometry.org": "Slow response observed in latest audit; security-header gaps.",
}

def add_line(doc, label):
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(" ______________________________________________________________")

def build(site, observation):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.5)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Guardian Customer Validation Form")
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(site).bold = True
    doc.add_paragraph("Please complete this short review so we can confirm whether Guardian's website-health results are accurate, understandable, and useful. This is a read-only review: Guardian will not change your website.")
    doc.add_heading("Owner and review details", level=1)
    for label in ["Business name", "Reviewer name and role", "Review date", "Permission to review confirmed yes or no"]:
        add_line(doc, label)
    doc.add_heading("What Guardian reported", level=1)
    doc.add_paragraph("The following is a technical observation from a bounded check. It is not a confirmed business problem until you review it.")
    p = doc.add_paragraph(); p.add_run("Observation: ").bold = True; p.add_run(observation)
    doc.add_heading("Your review", level=1)
    questions = [
        "Is this observation accurate for your website?  Yes / Partly / No / Not sure",
        "Was the explanation easy to understand?  Yes / No",
        "Does it affect your business or customers?  Yes / No / Not sure",
        "Did you take any action after seeing it?  Yes / No",
        "If action was taken, what did you change?",
        "Was the issue prevented or resolved?  Yes / No / Unknown",
        "What evidence supports your answer?",
        "How useful was this result?  1  2  3  4  5",
        "Would you continue using Guardian?  Yes / No / Unsure",
        "What was the most valuable result?",
        "What is the most important missing capability?",
        "May Guardian quote your feedback publicly without identifying private details?  Yes / No",
    ]
    for q in questions:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(q)
        if q.endswith(":"):
            doc.add_paragraph("______________________________________________________________")
    doc.add_heading("Return", level=1)
    doc.add_paragraph("Please return this completed form to the Guardian project owner. Your responses will be recorded separately from Guardian's technical results and will only be used for product validation unless you grant permission to quote them.")
    path = OUT / (site.replace(".", "-") + "-validation.docx")
    doc.save(path)

OUT.mkdir(parents=True, exist_ok=True)
for site, observation in SITES.items():
    build(site, observation)
print(f"Created {len(SITES)} documents in {OUT}")
